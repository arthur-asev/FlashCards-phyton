from io import BytesIO

import pytest
import xlwt
from fastapi.testclient import TestClient
from openpyxl import Workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app.api.v1.routes import files as file_routes
from app.db.base import Base
from app.db.session import get_db
from app.importers import spreadsheets
from app.main import app


@pytest.fixture
def client(tmp_path, monkeypatch) -> TestClient:
    engine = create_engine(
        "sqlite://",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    Base.metadata.create_all(engine)
    session_factory = sessionmaker(bind=engine, expire_on_commit=False)

    def override_get_db():
        with session_factory() as session:
            yield session

    app.dependency_overrides[get_db] = override_get_db
    monkeypatch.setattr(file_routes, "STORAGE_ROOT", tmp_path)
    with TestClient(app) as test_client:
        yield test_client
    app.dependency_overrides.clear()
    engine.dispose()


def make_xlsx() -> bytes:
    workbook = Workbook()
    workbook.active.title = "Cards"
    workbook.active.append(["Pergunta", "Resposta"])
    workbook.active.append(["  2 + 2?  ", "4"])
    workbook.create_sheet("Ignored").append(["Header"])
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def make_xls() -> bytes:
    workbook = xlwt.Workbook()
    sheet = workbook.add_sheet("Cards")
    sheet.write(0, 0, "Pergunta")
    sheet.write(0, 1, "Resposta")
    sheet.write(1, 0, "Capital do Brasil?")
    sheet.write(1, 1, "Brasília")
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def upload(client: TestClient, filename: str, content: bytes) -> str:
    response = client.post(
        "/api/v1/files/upload",
        files={"file": (filename, content, "application/octet-stream")},
    )
    assert response.status_code == 201, response.text
    return response.json()["id"]


def test_csv_preview_detects_delimiter_and_caps_rows(client: TestClient) -> None:
    file_id = upload(
        client,
        "cards.csv",
        b"Pergunta;Resposta\r\n  2 + 2? ; 4 \r\nCapital?;Brasilia\r\n",
    )

    response = client.get(f"/api/v1/files/{file_id}/preview?limit=1")

    assert response.status_code == 200
    preview = response.json()["preview"]
    assert preview["sheet_names"] == ["CSV"]
    assert preview["columns"] == ["Pergunta", "Resposta"]
    assert preview["rows"] == [["2 + 2?", "4"]]
    assert preview["row_count"] == 2
    assert preview["has_more"] is True


def test_xlsx_preview_selects_sheet_and_normalizes_cells(client: TestClient) -> None:
    file_id = upload(client, "cards.xlsx", make_xlsx())

    response = client.get(f"/api/v1/files/{file_id}/preview?sheet_name=Cards")

    assert response.status_code == 200
    preview = response.json()["preview"]
    assert preview["sheet_names"] == ["Cards", "Ignored"]
    assert preview["selected_sheet"] == "Cards"
    assert preview["rows"] == [["2 + 2?", "4"]]


def test_xls_preview_reads_legacy_workbook(client: TestClient) -> None:
    file_id = upload(client, "cards.xls", make_xls())

    response = client.get(f"/api/v1/files/{file_id}/preview")

    assert response.status_code == 200
    preview = response.json()["preview"]
    assert preview["columns"] == ["Pergunta", "Resposta"]
    assert preview["rows"] == [["Capital do Brasil?", "Brasília"]]


def test_csv_mapping_returns_normalized_valid_rows(client: TestClient) -> None:
    file_id = upload(
        client,
        "cards.csv",
        "Pergunta,Resposta,Matéria\n 2 + 2? , 4 , Matemática \n".encode("utf-8"),
    )

    response = client.post(
        f"/api/v1/files/{file_id}/validate",
        json={
            "mapping": {
                "front": "Pergunta",
                "back": "Resposta",
                "subject": "Mat\u00e9ria",
            }
        },
    )

    assert response.status_code == 200
    result = response.json()
    assert result["valid"] is True
    assert result["valid_rows"] == 1
    assert result["normalized_rows"] == [
        {"front": "2 + 2?", "back": "4", "subject": "Matemática"}
    ]


def test_validation_reports_duplicates_empty_and_missing_values(
    client: TestClient,
) -> None:
    file_id = upload(
        client,
        "cards.csv",
        b"front,back\nQuestion,Answer\n question , answer \nOnly front,\n,\n",
    )

    response = client.post(
        f"/api/v1/files/{file_id}/validate",
        json={"mapping": {"front": "front", "back": "back"}},
    )

    assert response.status_code == 200
    result = response.json()
    assert result["valid"] is False
    assert {issue["code"] for issue in result["issues"]} == {
        "DUPLICATE_CARD",
        "EMPTY_ROW",
        "REQUIRED_VALUE_MISSING",
    }


def test_validation_reports_unmapped_columns(client: TestClient) -> None:
    file_id = upload(client, "cards.csv", b"Question,Answer\nFront,Back\n")

    response = client.post(
        f"/api/v1/files/{file_id}/validate",
        json={"mapping": {"front": "Question", "back": "Missing"}},
    )

    assert response.status_code == 200
    assert "COLUMN_NOT_FOUND" in {issue["code"] for issue in response.json()["issues"]}


def test_validation_rejects_workbook_without_data_rows(client: TestClient) -> None:
    file_id = upload(client, "cards.csv", b"front,back\n")

    response = client.post(
        f"/api/v1/files/{file_id}/validate",
        json={"mapping": {"front": "front", "back": "back"}},
    )

    assert response.status_code == 200
    result = response.json()
    assert result["valid"] is False
    assert result["issues"][0]["code"] == "NO_DATA_ROWS"


def test_validation_reports_irregular_row_and_excessive_cell(
    client: TestClient,
) -> None:
    too_long_front = "Q" * (spreadsheets.MAX_CELL_CHARACTERS + 1)
    content = f"front,back\nA,B,extra\n{too_long_front},B\n".encode("utf-8")
    file_id = upload(client, "cards.csv", content)

    response = client.post(
        f"/api/v1/files/{file_id}/validate",
        json={"mapping": {"front": "front", "back": "back"}},
    )

    assert response.status_code == 200
    result = response.json()
    assert result["valid"] is False
    assert {issue["code"] for issue in result["issues"]} == {
        "COLUMN_COUNT_MISMATCH",
        "CONTENT_TOO_LONG",
    }


def test_preview_rejects_unknown_sheet_and_row_limit(
    client: TestClient, monkeypatch
) -> None:
    file_id = upload(client, "cards.csv", b"front,back\nA,B\nC,D\n")
    missing_sheet = client.get(f"/api/v1/files/{file_id}/preview?sheet_name=Missing")
    assert missing_sheet.status_code == 422

    monkeypatch.setattr(spreadsheets, "MAX_SPREADSHEET_ROWS", 1)
    too_many_rows = client.get(f"/api/v1/files/{file_id}/preview")
    assert too_many_rows.status_code == 413
