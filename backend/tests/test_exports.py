import csv
import json
from io import BytesIO, StringIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.main import app

client = TestClient(app)
CARDS = [
    {
        "front": "O que é fotossíntese?",
        "back": "Conversão de energia luminosa em energia química.",
        "explanation": "Ocorre nos cloroplastos.",
        "example": "Plantas usam luz solar.",
        "difficulty": "easy",
        "tags": ["biologia", "energia"],
        "subject": "Biologia",
        "topic": "Botânica",
        "source": "Caderno 1",
    }
]


@pytest.mark.parametrize(
    ("export_format", "expected_media_type", "expected_extension"),
    [
        ("csv", "text/csv; charset=utf-8", "csv"),
        ("json", "application/json", "json"),
        (
            "xlsx",
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            "xlsx",
        ),
    ],
)
def test_export_response_headers(
    export_format: str,
    expected_media_type: str,
    expected_extension: str,
) -> None:
    response = client.post(f"/api/v1/export/{export_format}", json={"cards": CARDS})

    assert response.status_code == 200
    assert response.headers["content-type"] == expected_media_type
    assert response.headers["content-disposition"] == (
        f'attachment; filename="flashcards.{expected_extension}"'
    )
    assert response.content


def test_csv_export_uses_utf8_and_escapes_delimiters() -> None:
    payload = [
        {**CARDS[0], "front": "Frente, com vírgula", "back": 'Resposta "citada"'}
    ]

    response = client.post("/api/v1/export/csv", json={"cards": payload})
    rows = list(csv.reader(StringIO(response.content.decode("utf-8"))))

    assert rows[0] == [
        "front",
        "back",
        "explanation",
        "example",
        "difficulty",
        "tags",
        "subject",
        "topic",
        "source",
    ]
    assert rows[1][0] == "Frente, com vírgula"
    assert rows[1][1] == 'Resposta "citada"'
    assert rows[1][5] == "biologia; energia"
    assert "vírgula".encode("utf-8") in response.content


def test_json_export_preserves_unicode_and_tag_array() -> None:
    response = client.post("/api/v1/export/json", json={"cards": CARDS})

    assert json.loads(response.content) == CARDS
    assert "fotossíntese".encode("utf-8") in response.content


def test_xlsx_export_contains_all_fields() -> None:
    response = client.post("/api/v1/export/xlsx", json={"cards": CARDS})
    workbook = load_workbook(BytesIO(response.content), read_only=True, data_only=True)
    worksheet = workbook["Flashcards"]
    rows = list(worksheet.iter_rows(values_only=True))
    workbook.close()

    assert rows[0] == (
        "front",
        "back",
        "explanation",
        "example",
        "difficulty",
        "tags",
        "subject",
        "topic",
        "source",
    )
    assert rows[1][0] == CARDS[0]["front"]
    assert rows[1][5] == "biologia; energia"


@pytest.mark.parametrize("export_format", ["csv", "json", "xlsx"])
def test_empty_export_is_supported(export_format: str) -> None:
    response = client.post(f"/api/v1/export/{export_format}", json={"cards": []})

    assert response.status_code == 200
    assert response.content


def test_export_rejects_missing_required_card_fields() -> None:
    response = client.post(
        "/api/v1/export/json", json={"cards": [{"front": "Only front"}]}
    )

    assert response.status_code == 422


def test_export_rejects_unknown_fields() -> None:
    response = client.post(
        "/api/v1/export/json",
        json={"cards": [{"front": "Question", "back": "Answer", "unknown": True}]},
    )

    assert response.status_code == 422
