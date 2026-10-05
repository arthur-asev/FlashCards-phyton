import csv
from collections.abc import Iterable
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from zipfile import BadZipFile

import xlrd
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from xlrd.biffh import XLRDError

from app.importers.base import Importer

MAX_SPREADSHEET_ROWS = 10000
MAX_CELL_CHARACTERS = 10000
ALLOWED_FIELDS = {
    "front",
    "back",
    "explanation",
    "subject",
    "topic",
    "tags",
    "difficulty",
    "source",
}
REQUIRED_FIELDS = {"front", "back"}


class SpreadsheetError(ValueError):
    """Raised when a spreadsheet cannot be read or mapped safely."""


class SpreadsheetTooLargeError(SpreadsheetError):
    """Raised when a spreadsheet exceeds the synchronous row limit."""


@dataclass(frozen=True)
class SpreadsheetTable:
    sheet_names: list[str]
    selected_sheet: str
    headers: list[str]
    rows: list[list[str]]


def normalize_cell(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _collect_rows(rows: Iterable[Iterable[object]]) -> list[list[str]]:
    collected: list[list[str]] = []
    for row in rows:
        if len(collected) >= MAX_SPREADSHEET_ROWS:
            raise SpreadsheetTooLargeError(
                f"Spreadsheets are limited to {MAX_SPREADSHEET_ROWS} data rows per request."
            )
        collected.append([normalize_cell(value) for value in row])
    return collected


def _read_csv(path: Path) -> SpreadsheetTable:
    try:
        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            sample = handle.read(4096)
            handle.seek(0)
            try:
                dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
            except csv.Error:
                dialect = csv.excel
            reader = csv.reader(handle, dialect)
            raw_headers = next(reader, None)
            if raw_headers is None:
                raise SpreadsheetError("The spreadsheet has no header row.")
            headers = [normalize_cell(value) for value in raw_headers]
            if not any(headers):
                raise SpreadsheetError("The header row has no column names.")
            rows = _collect_rows(reader)
    except UnicodeDecodeError as exc:
        raise SpreadsheetError("CSV files must use UTF-8 encoding.") from exc
    except csv.Error as exc:
        raise SpreadsheetError("The CSV file has invalid row formatting.") from exc

    return SpreadsheetTable(["CSV"], "CSV", headers, rows)


def _select_sheet(sheet_names: list[str], sheet_name: str | None) -> str:
    if not sheet_names:
        raise SpreadsheetError("The workbook contains no worksheets.")
    selected = sheet_name or sheet_names[0]
    if selected not in sheet_names:
        raise SpreadsheetError(f"Worksheet '{selected}' was not found.")
    return selected


def _read_xlsx(path: Path, sheet_name: str | None) -> SpreadsheetTable:
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except (InvalidFileException, BadZipFile, OSError) as exc:
        raise SpreadsheetError("The XLSX file is invalid or unreadable.") from exc

    try:
        sheet_names = workbook.sheetnames
        selected = _select_sheet(sheet_names, sheet_name)
        sheet = workbook[selected]
        iterator = sheet.iter_rows(values_only=True)
        raw_headers = next(iterator, None)
        if raw_headers is None:
            raise SpreadsheetError("The selected worksheet has no header row.")
        headers = [normalize_cell(value) for value in raw_headers]
        if not any(headers):
            raise SpreadsheetError("The header row has no column names.")
        rows = _collect_rows(iterator)
    finally:
        workbook.close()

    return SpreadsheetTable(sheet_names, selected, headers, rows)


def _read_xls(path: Path, sheet_name: str | None) -> SpreadsheetTable:
    try:
        workbook = xlrd.open_workbook(filename=str(path), on_demand=True)
    except (XLRDError, OSError) as exc:
        raise SpreadsheetError("The XLS file is invalid or unreadable.") from exc

    try:
        sheet_names = workbook.sheet_names()
        selected = _select_sheet(sheet_names, sheet_name)
        sheet = workbook.sheet_by_name(selected)
        if sheet.nrows == 0:
            raise SpreadsheetError("The selected worksheet has no header row.")
        headers = [normalize_cell(value) for value in sheet.row_values(0)]
        if not any(headers):
            raise SpreadsheetError("The header row has no column names.")
        rows: list[list[str]] = []
        for row_index in range(1, sheet.nrows):
            if len(rows) >= MAX_SPREADSHEET_ROWS:
                raise SpreadsheetTooLargeError(
                    f"Spreadsheets are limited to {MAX_SPREADSHEET_ROWS} data rows per request."
                )
            values: list[str] = []
            for column_index in range(sheet.ncols):
                cell = sheet.cell(row_index, column_index)
                value = cell.value
                if cell.ctype == xlrd.XL_CELL_DATE:
                    value = xlrd.xldate_as_datetime(value, workbook.datemode)
                values.append(normalize_cell(value))
            rows.append(values)
    finally:
        workbook.release_resources()

    return SpreadsheetTable(sheet_names, selected, headers, rows)


def read_spreadsheet(
    path: str | Path, sheet_name: str | None = None
) -> SpreadsheetTable:
    file_path = Path(path)
    extension = file_path.suffix.lower()
    if extension == ".csv":
        if sheet_name not in (None, "CSV"):
            raise SpreadsheetError(f"Worksheet '{sheet_name}' was not found.")
        return _read_csv(file_path)
    if extension == ".xlsx":
        return _read_xlsx(file_path, sheet_name)
    if extension == ".xls":
        return _read_xls(file_path, sheet_name)
    raise SpreadsheetError("Only CSV, XLSX and XLS spreadsheets are supported.")


class SpreadsheetImporter(Importer[SpreadsheetTable]):
    def read(self, path: str, sheet_name: str | None = None) -> SpreadsheetTable:
        return read_spreadsheet(path, sheet_name)


def validate_column_mapping(
    table: SpreadsheetTable,
    mapping: dict[str, str | int],
) -> dict[str, object]:
    issues: list[dict[str, int | str | None]] = []
    normalized_headers = [header.casefold() for header in table.headers]

    for field in mapping:
        if field not in ALLOWED_FIELDS:
            issues.append(
                {
                    "row": None,
                    "code": "UNKNOWN_FIELD",
                    "message": f"Unknown target field: {field}.",
                }
            )
    for field in sorted(REQUIRED_FIELDS - mapping.keys()):
        issues.append(
            {
                "row": None,
                "code": "REQUIRED_COLUMN_NOT_MAPPED",
                "message": f"A source column must be mapped to '{field}'.",
            }
        )
    uses_header_names = any(isinstance(source, str) for source in mapping.values())
    if uses_header_names and len(normalized_headers) != len(set(normalized_headers)):
        issues.append(
            {
                "row": 1,
                "code": "DUPLICATE_HEADER",
                "message": "Column headers must be unique.",
            }
        )
    if not table.rows:
        issues.append(
            {
                "row": None,
                "code": "NO_DATA_ROWS",
                "message": "The spreadsheet has no data rows.",
            }
        )

    column_indexes: dict[str, int] = {}
    for field, source_column in mapping.items():
        if isinstance(source_column, int) and not isinstance(source_column, bool):
            if not 0 <= source_column < len(table.headers):
                issues.append(
                    {
                        "row": None,
                        "code": "COLUMN_INDEX_OUT_OF_RANGE",
                        "message": f"Column index {source_column} is outside the spreadsheet.",
                    }
                )
            else:
                column_indexes[field] = source_column
        elif source_column not in table.headers:
            issues.append(
                {
                    "row": None,
                    "code": "COLUMN_NOT_FOUND",
                    "message": f"Source column '{source_column}' was not found.",
                }
            )
        else:
            column_indexes[field] = table.headers.index(source_column)

    if issues:
        return {
            "valid": False,
            "checked_rows": 0,
            "valid_rows": 0,
            "issues": issues,
            "normalized_rows": [],
        }

    normalized_rows: list[dict[str, str]] = []
    seen_cards: set[tuple[str, str]] = set()
    valid_rows = 0

    for row_number, row in enumerate(table.rows, start=2):
        if not any(value.strip() for value in row):
            issues.append(
                {"row": row_number, "code": "EMPTY_ROW", "message": "The row is empty."}
            )
            continue
        row_is_valid = len(row) == len(table.headers)
        if not row_is_valid:
            issues.append(
                {
                    "row": row_number,
                    "code": "COLUMN_COUNT_MISMATCH",
                    "message": "The row has a different number of cells than the header.",
                }
            )

        normalized = {
            field: row[column_index].strip() if column_index < len(row) else ""
            for field, column_index in column_indexes.items()
        }
        for field in sorted(REQUIRED_FIELDS):
            if not normalized.get(field):
                issues.append(
                    {
                        "row": row_number,
                        "code": "REQUIRED_VALUE_MISSING",
                        "message": f"The '{field}' value is required.",
                    }
                )
                row_is_valid = False

        for field, value in normalized.items():
            if len(value) > MAX_CELL_CHARACTERS:
                issues.append(
                    {
                        "row": row_number,
                        "code": "CONTENT_TOO_LONG",
                        "message": f"The '{field}' value exceeds {MAX_CELL_CHARACTERS} characters.",
                    }
                )
                row_is_valid = False

        if normalized.get("front") and normalized.get("back"):
            card_key = (normalized["front"].casefold(), normalized["back"].casefold())
            if card_key in seen_cards:
                issues.append(
                    {
                        "row": row_number,
                        "code": "DUPLICATE_CARD",
                        "message": "The card is duplicated.",
                    }
                )
                row_is_valid = False
            else:
                seen_cards.add(card_key)

        normalized_rows.append(normalized)
        valid_rows += int(row_is_valid)

    return {
        "valid": not issues,
        "checked_rows": len(table.rows),
        "valid_rows": valid_rows,
        "issues": issues,
        "normalized_rows": normalized_rows,
    }
